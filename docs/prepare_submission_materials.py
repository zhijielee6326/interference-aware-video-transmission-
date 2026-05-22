from __future__ import annotations

import re
import shutil
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(r"C:\Users\m1889\Desktop\毕设")
OUT = ROOT / "提交材料_李智杰"
MAIN_DOCX = ROOT / "毕业论文_李智杰.docx"
DESKTOP_DRAFT_DOCX = Path(r"C:\Users\m1889\Desktop\毕业论文_李智杰.docx")
FOREIGN_PDF = ROOT / "参考文献" / "fncom-17-1309694.pdf"

FINAL_DOCX = OUT / "毕业论文_李智杰_终稿.docx"
CHECK_DOCX = OUT / "毕业论文_李智杰_查重稿.docx"
TRANSLATION_DOCX = OUT / "外文文献译文_李智杰.docx"
APPENDIX_INDEX_DOCX = OUT / "附录目录与说明_李智杰.docx"
INSPECTION_INDEX_DOCX = OUT / "抽检材料清单_李智杰.docx"
ALIGNMENT_REPORT_DOCX = OUT / "材料核对报告_李智杰.docx"

TITLE = "面向信号干扰检测与识别的视频传输平台开发"
EN_TITLE = "Interference-Aware Adaptive Video Transmission Platform Based on CCNN"
COLLEGE = "信息与通信工程学院"
MAJOR = "通信工程（卓越）"
STUDENT = "李智杰"
CLASS_ID = "通信2201 / 2022010843"
ADVISOR = "刘磊"
DATE_TEXT = "2026年5月20日"


CN_ABSTRACT = [
    (
        "随着5G/6G通信系统向高频段、大带宽和高密度连接方向发展，复杂电磁环境中的同频、邻频及恶意干扰会显著降低视频业务链路质量。"
        "为提升通信系统对干扰态势的感知能力和视频传输可靠性，本文设计并实现了一套面向信号干扰检测与识别的自适应视频传输平台。"
        "平台以USRP B210软件无线电为硬件基础，构建了覆盖视频采集、JPEG编码、帧结构分包、前向纠错、QPSK调制解调、信道传输、干扰识别和传输策略调整的完整链路。"
    ),
    (
        "在干扰识别方面，本文构建了复合卷积神经网络CCNN模型，融合卷积特征提取、残差学习、SE通道注意力和多头因果注意力机制，"
        "实现对扫频、多音、窄带调幅、窄带调频、单音、正弦波及正常信号共7类样本的识别。实验结果表明，模型在JSR不低于0 dB时识别准确率达到99.9%，"
        "在JSR=-2 dB条件下仍保持较高鲁棒性。系统测试结果显示，平台能够根据识别结果动态调整JPEG质量、FEC冗余和帧率，完成仿真模式与USRP实机模式下的闭环验证。"
    ),
]
CN_KEYWORDS = "关键词：干扰识别；卷积神经网络；自适应传输；软件无线电；USRP B210"

EN_ABSTRACT = [
    (
        "With the evolution of 5G and 6G communication systems toward wider bandwidth, higher carrier frequency, and denser network deployment, "
        "co-channel, adjacent-channel, and intentional jamming may seriously degrade the quality of video transmission links. "
        "To improve interference awareness and transmission reliability in complex electromagnetic environments, this thesis designs and implements "
        "an adaptive video transmission platform for signal interference detection and recognition. Based on the USRP B210 software-defined radio, "
        "the platform integrates video capture, JPEG encoding, packet framing, forward error correction, QPSK modulation and demodulation, "
        "channel transmission, interference recognition, and adaptive parameter control into a complete end-to-end chain."
    ),
    (
        "For interference recognition, a composite convolutional neural network, named CCNN, is constructed by combining convolutional feature extraction, "
        "residual learning, squeeze-and-excitation channel attention, and multi-head causal attention. The model classifies seven categories, including "
        "linear frequency modulation jamming, multi-tone jamming, narrowband amplitude modulation jamming, narrowband frequency modulation jamming, "
        "single-tone jamming, sinusoidal jamming, and clean signals. Experimental results show that the recognition accuracy reaches 99.9% when JSR is "
        "not lower than 0 dB and remains robust at JSR = -2 dB. System tests further demonstrate that the proposed platform can dynamically adjust "
        "JPEG quality, FEC redundancy, and frame rate according to recognition results, and can complete closed-loop verification in both simulation "
        "and USRP hardware-in-the-loop modes."
    ),
]
EN_KEYWORDS = "Keywords: Interference Recognition; Convolutional Neural Network; Adaptive Transmission; Software-Defined Radio; USRP B210"


EXPANSIONS: dict[str, list[str]] = {
    "研究背景与意义": [
        (
            "从工程应用角度看，视频业务链路不仅要求较高的数据吞吐量，还要求端到端时延、丢包率和画面连续性同时满足约束。"
            "当通信链路受到窄带或扫频干扰时，接收端首先表现为信噪比下降、同步序列检测不稳定和误码率上升，随后才体现为视频画面卡顿、马赛克或中断。"
            "因此，若系统只能在视频质量已经明显下降后再被动调参，往往难以及时恢复链路质量。本文将干扰检测识别前置到物理层和基带信号处理阶段，"
            "通过对IQ样本和功率谱特征的分析提前获取信道异常信息，再把识别结果反馈给视频传输控制模块，为上层编码质量和冗余策略调整提供依据。"
        ),
        (
            "与传统固定参数的视频传输方案相比，干扰感知型传输平台更强调跨层协同：物理层负责采集和还原信号状态，算法层负责判断干扰类型和可信度，"
            "应用层则根据链路状态动态改变视频编码质量、FEC冗余和帧率。这样的设计能够在不同干扰强度下取得可靠性与实时性的折中。"
            "在无干扰或轻度干扰场景中，系统优先保证较高帧率和画面清晰度；在强干扰场景中，系统牺牲部分视觉质量以提升解码成功率，"
            "从而避免视频链路完全中断。这种策略符合智能通信系统由静态配置向环境自适应演进的发展趋势。"
        ),
    ],
    "干扰信号检测与识别": [
        (
            "现有干扰识别研究大致可分为基于人工特征的方法和基于深度学习的方法。前者通常提取功率谱峰值、谱熵、频率重心、带宽占用率等统计特征，"
            "再结合阈值判决、支持向量机或浅层分类器完成识别，其优点是计算量较小、可解释性较强，但对复杂时变干扰和低JSR样本的泛化能力有限。"
            "后者直接从IQ序列、频谱图或时频图中学习特征表示，能够减少人工建模对先验知识的依赖，尤其适合多类型干扰混合、信道参数不确定的场景。"
            "本文采用以深度网络为核心、以功率谱特征为辅助的方案，既保留传统方法对能量变化敏感的优势，又利用神经网络增强类别区分能力。"
        ),
        (
            "在网络结构选择上，单纯CNN擅长提取局部波形和频谱纹理，但对长序列范围内的上下文依赖建模不足；循环网络具有时序表达能力，"
            "但训练效率和并行计算能力相对受限；Transformer类结构能够刻画长距离依赖，却需要解决一维信号局部细节提取不足的问题。"
            "基于上述考虑，本文的CCNN模型采用复合结构，将卷积、残差连接、通道注意力和因果注意力组合起来，使模型同时具备局部细节感知、"
            "梯度稳定传播、通道重要性重标定和时序上下文建模能力。该结构与论文任务中的实时干扰识别需求具有较好匹配性。"
        ),
    ],
    "软件无线电与USRP B210": [
        (
            "USRP B210能够在较宽频率范围内完成射频信号收发，并通过UHD驱动向上层软件提供统一接口，适合用于通信算法从仿真到实机的迁移验证。"
            "本文选择USRP B210作为硬件平台，主要考虑其双通道收发能力、较高采样率、成熟的软件生态和较低的实验门槛。"
            "在系统实现中，USRP负责完成基带IQ数据与射频信号之间的转换，Python程序负责帧封装、调制解调、干扰注入、模型推理和可视化显示。"
            "这种软硬件分工既便于快速调试算法，也便于根据实验需要调整采样率、中心频率、增益和缓冲区等关键参数。"
        ),
    ],
    "QPSK调制解调": [
        (
            "QPSK调制方式在本系统中承担着连接视频字节流与无线基带信号的关键作用。发送端将经过分包和冗余编码后的比特序列按照两比特一组映射到四个相位点，"
            "再通过根升余弦滤波器进行脉冲成形，以降低带外泄漏和符号间干扰。接收端则需要在存在噪声、频偏和幅度变化的情况下完成符号判决。"
            "由于本文重点是构建可验证的干扰识别与自适应传输平台，QPSK在复杂度、实现难度和实验可解释性之间取得了较好平衡：它比BPSK具有更高频谱效率，"
            "又比高阶QAM对信噪比和同步精度的要求更低，适合作为毕业设计阶段软硬件闭环验证的基础调制方案。"
        ),
        (
            "在实机链路中，QPSK解调性能受到多个因素共同影响。载波频偏会导致星座图整体旋转，采样时偏会使最佳判决点偏离符号中心，"
            "自动增益控制不稳定则会改变星座点幅度分布。本文在接收端加入AGC、前导序列同步和频偏估计校正，目的不是追求通信体制的完整商用化实现，"
            "而是保证平台能够在可控实验条件下稳定传递视频分片，并为后续干扰识别模块提供足够可靠的信号输入。"
            "这一设计使系统既能展示完整数字通信链路，也能清楚反映干扰对不同处理环节的影响。"
        ),
    ],
    "系统架构": [
        (
            "系统总体架构遵循模块化和可替换原则，按照数据流方向划分为发送端、信道/干扰环境、接收端、识别决策模块和可视化模块。"
            "发送端负责从摄像头或测试视频源获取图像帧，并根据当前策略完成JPEG压缩、分片、添加帧头和CRC校验，再经过FEC冗余编码和QPSK调制生成基带信号。"
            "接收端完成采样数据缓存、帧同步、频偏估计、AGC归一化、QPSK解调、FEC恢复、CRC检查和JPEG解码。"
            "识别决策模块从接收侧信号中提取用于CCNN推理的片段，并将干扰类别、置信度和干扰等级反馈给发送控制策略。"
        ),
        (
            "为了保证系统在仿真和实机两种模式下具有一致的软件接口，本文将信道抽象为统一的数据交换层。"
            "在仿真模式中，该层由SimulatedChannel完成，可直接注入不同类型和强度的干扰信号，便于重复实验和参数扫描；"
            "在实机模式中，该层由USRP收发链路替代，通过UHD接口完成真实射频采样。上层模块不直接依赖具体信道实现，"
            "因此同一套TX/RX逻辑可以在两种模式下运行。这种设计降低了调试复杂度，也使实验结果能够更清晰地反映算法本身和硬件链路之间的差异。"
        ),
    ],
    "系统参数配置": [
        (
            "系统参数配置直接影响平台在不同测试场景下的稳定性和可复现实验能力。本文将采样率、中心频率、符号率、JPEG基准质量、FEC冗余倍数、"
            "干扰注入强度和模型推理窗口长度等关键参数集中管理，避免这些参数分散在多个模块中造成调试困难。集中配置还有助于保证论文实验与程序实现之间的一致性："
            "当第六章需要对比不同JSR条件下的性能时，只需调整统一配置项即可复现实验，而不需要修改底层算法代码。"
        ),
        (
            "参数配置同时体现了仿真模式和实机模式的差异。仿真模式不受硬件带宽、USB缓冲和射频前端非理想因素限制，适合使用较高分辨率和较高JPEG质量验证算法逻辑；"
            "实机模式需要考虑USRP采样率、主机处理能力和实时显示开销，因此采用较低视频分辨率和更保守的编码参数。"
            "这种差异不是系统能力不足，而是工程验证中必须面对的资源约束。通过在配置层显式区分两种模式，系统能够在算法验证和硬件演示之间保持清晰边界。"
        ),
    ],
    "网络总体架构": [
        (
            "CCNN网络的输入为定长IQ复数序列，经预处理后转化为适合一维卷积处理的张量。第一阶段卷积层用于提取局部幅相变化、频谱边缘和周期性扰动等低层特征；"
            "随后残差模块通过跳连结构缓解深层网络训练中的梯度衰减问题，使模型可以在保持较小参数量的同时学习更抽象的干扰模式。"
            "SE通道注意力模块对不同特征通道进行权重重标定，使模型在不同JSR条件下能够自动突出对分类更有贡献的频率纹理或时序片段。"
        ),
        (
            "多头因果注意力模块用于进一步建模信号序列内部的长距离依赖关系。对于扫频干扰和正弦波干扰而言，干扰特征往往并不集中在单个采样点附近，"
            "而是表现为跨时间窗口的连续变化；对于多音和窄带干扰而言，局部谱峰结构和全局能量分布也需要联合判断。"
            "多头机制能够让不同注意力头关注不同时间尺度和特征子空间，因果约束则避免模型在实时推理场景中过度依赖未来信息，"
            "从而更贴近在线检测任务的使用方式。"
        ),
    ],
    "训练数据生成": [
        (
            "训练数据质量决定了干扰识别模型的上限。本文围绕六类典型干扰和正常信号构造样本，覆盖不同JSR、频率偏移、相位和噪声条件，"
            "使模型在训练阶段能够看到足够多的变化形式。与只在单一信噪比下生成样本相比，宽JSR范围的数据集更接近真实通信环境，"
            "能够降低模型对某一固定功率条件的过拟合风险。数据生成过程还保留类别编号、JSR参数和样本划分信息，便于后续复现实验和定位误判来源。"
        ),
        (
            "本文同时比较了MATLAB和Python两种数据生成方式。MATLAB脚本便于快速验证单类干扰的数学模型和波形特征，但多脚本结构在批量生成、系统集成和在线仿真方面不够灵活。"
            "Python生成器将多类干扰逻辑整合到统一接口中，能够按需生成训练样本，也能够与视频传输仿真信道直接连接。"
            "因此，Python方案不仅减少了中间文件存储压力，也使模型训练、推理测试和系统级实验形成更连贯的工作流。"
        ),
    ],
    "TX端设计": [
        (
            "发送端的设计重点在于把视频帧转化为可在受扰信道中稳定传输的数据单元。首先，系统对输入视频帧进行尺寸调整和JPEG压缩，"
            "通过控制压缩质量在清晰度和码率之间建立可调节接口。随后，压缩后的字节流被切分为固定长度载荷，每个分片附加同步字、帧序号、分片序号、"
            "总分片数和CRC校验字段，使接收端能够在存在丢包和误判的情况下恢复帧边界。分片级设计降低了单次错误对整帧视频的影响，也便于FEC冗余机制针对包级丢失进行恢复。"
        ),
        (
            "在调制环节，系统采用QPSK作为基础调制方式。QPSK每个符号携带2比特信息，复杂度适中，便于在Python环境中实现成形滤波、符号映射和解调判决。"
            "发送端在调制后引入根升余弦脉冲成形，以降低码间串扰并控制信号带宽。通过统一的帧结构和调制接口，系统能够在后续研究中进一步扩展到BPSK、"
            "16QAM或自适应调制编码方案，而无需重写上层视频封装逻辑。"
        ),
    ],
    "RX端设计": [
        (
            "接收端是系统闭环验证中最容易暴露工程问题的部分。理论上，接收端只需完成解调、解码和视频重组，但在实际采样数据流中，"
            "帧头可能被噪声误触发，分片可能乱序或丢失，CRC校验也可能因单个比特错误而失败。为此，本文在RX端采用分层处理流程："
            "底层负责持续读取IQ数据并完成同步和解调，中间层根据帧结构解析分片并校验完整性，上层负责FEC恢复、JPEG解码和图像显示。"
            "这种分层设计使每个环节的输入输出更加明确，便于在仿真模式和USRP实机模式下分别定位问题。"
        ),
        (
            "在视频重组过程中，系统需要处理实时性和完整性之间的矛盾。如果接收端等待所有分片完全到齐，视频延迟会随丢包而快速增加；"
            "如果过早解码，又可能导致图像损坏。本文通过分片序号、总分片数和CRC校验判断帧完整性，并结合FEC冗余尽可能恢复缺失分片。"
            "对于无法恢复的帧，系统倾向于丢弃该帧并继续处理后续数据，以保持视频播放连续性。该策略符合实时视频业务的特点：在干扰环境下，"
            "短时画质下降通常比长时间停顿更容易被接受。"
        ),
    ],
    "自适应传输策略": [
        (
            "自适应传输策略的核心是避免识别结果的瞬时波动直接导致传输参数频繁跳变。实际链路中，模型置信度可能随噪声、同步误差和短时信道衰落而波动，"
            "若每一次分类输出都立即触发参数切换，视频质量会出现明显抖动。本文引入多级滞后计数器机制，对干扰出现、干扰等级变化和干扰消失分别设置确认门限。"
            "只有连续多次观测满足条件时，系统才改变JPEG质量、FEC倍数和帧率，从而提高策略稳定性。"
        ),
        (
            "策略映射关系按照干扰对链路破坏程度进行分级。无干扰时系统采用较高JPEG质量和较高帧率，减少冗余以提升有效吞吐；"
            "轻度干扰时适当降低图像质量并提高FEC冗余，使接收端具备更强的分片恢复能力；强干扰时进一步降低帧率和码率，"
            "把有限信道资源优先用于保证关键帧能够被正确接收。该策略并不追求单一指标最优，而是在画面清晰度、实时性和可靠性之间实现工程可用的动态平衡。"
        ),
    ],
    "实验环境": [
        (
            "实验验证分为模型性能测试、仿真链路测试和USRP实机测试三个层次。模型性能测试主要关注CCNN在不同JSR条件和不同干扰类别下的识别准确率，"
            "用于评价算法本身的分类能力；仿真链路测试在可控信道中验证TX/RX模块、FEC恢复和自适应策略是否能够按预期工作；"
            "USRP实机测试则进一步引入真实射频链路中的频偏、采样时偏、增益变化和USB传输缓冲限制，用于检验系统从软件仿真迁移到硬件平台后的可用性。"
        ),
        (
            "这种分层测试方法能够减少问题定位难度。若模型离线准确率不足，说明需要改进数据生成或网络结构；若仿真链路不稳定，"
            "说明帧结构、同步、FEC或解码逻辑需要调整；若实机链路与仿真结果存在差距，则重点分析硬件增益、频偏估计、采样率匹配和缓冲区设置。"
            "本文通过逐层验证，使系统测试不仅给出结果数值，也能够说明各模块在最终平台中的作用边界和改进方向。"
        ),
    ],
    "整体分类性能": [
        (
            "从整体分类结果看，CCNN模型在中高JSR区间表现稳定，说明复合网络结构能够有效学习六类典型干扰与正常信号之间的差异。"
            "在低JSR区间，NAM、STJ等窄带干扰与正常QPSK信号在局部频谱形态上较为接近，因此更容易出现误判。"
            "这一现象与通信信号的物理特性一致：当干扰功率较低时，窄带干扰只在有限频点上改变能量分布，对整体波形的影响并不明显。"
            "因此，后续若要进一步提升低JSR识别性能，应重点增加弱干扰样本、引入时频联合特征，并考虑通过置信度校准降低误触发概率。"
        ),
    ],
    "与其他方法对比": [
        (
            "与传统基于功率谱阈值的检测方法相比，本文模型不仅能够判断是否存在干扰，还能够进一步区分干扰类型，为自适应传输策略提供更细粒度的输入。"
            "与单一CNN模型相比，CCNN通过残差结构和注意力机制增强了特征表达能力，在保持较小参数量的同时提升了复杂样本的分类稳定性。"
            "与需要大规模图像化时频变换的模型相比，本文方案直接面向IQ序列和轻量特征进行处理，便于集成到实时视频传输链路中。"
            "这种对比说明，模型设计不能只追求离线准确率，还需要综合考虑参数规模、推理延迟、工程实现难度和系统集成成本。"
        ),
    ],
    "本文工作总结": [
        (
            "总体来看，本文完成的系统不仅是一个单独的干扰识别模型，也是一套围绕视频业务构建的端到端验证平台。"
            "模型层面的识别结果能够被传输控制策略实际使用，传输策略的效果又可以通过接收帧率、解码成功率和可视化界面直观反馈。"
            "这种闭环设计使论文工作从离线分类准确率扩展到通信系统可用性验证，更符合复杂电磁环境下智能通信平台的工程需求。"
        ),
    ],
    "不足与展望": [
        (
            "尽管本文完成了从算法模型到视频传输平台的整体实现，但系统仍具有进一步优化空间。首先，当前平台主要采用应用层参数自适应，"
            "尚未将调制方式、编码码率、发射功率和频谱资源选择纳入统一控制。若后续引入自适应调制编码、跳频或扩频机制，"
            "系统将能够在物理层和应用层之间形成更完整的跨层抗干扰策略。其次，当前模型主要识别训练集中定义的典型干扰类型，"
            "对于未知干扰的拒识和增量学习能力仍需加强。未来可通过开放集识别、置信度校准和在线更新机制，提高系统面对新型干扰时的安全性和可扩展性。"
        ),
        (
            "此外，USRP实机模式中的频偏、采样时偏和缓冲区溢出问题说明，真实硬件链路比仿真环境更复杂。后续可以将部分高实时性处理迁移到更底层的C++、GNU Radio或FPGA模块中，"
            "降低Python解释执行和可视化刷新带来的延迟。也可以扩展为多节点实验平台，让多个接收节点协同感知干扰空间分布，"
            "从而实现更贴近实际网络的干扰定位和资源调度。上述改进将使本文平台从单链路验证原型进一步发展为可扩展的智能抗干扰通信实验系统。"
        ),
        (
            "最后，从工程管理角度看，本文已形成论文、程序、测试记录、外文译文和抽检支撑材料之间的对应关系。后续若继续完善该平台，"
            "应同步维护实验配置、模型版本、数据集来源和测试脚本说明，避免算法结果与论文描述脱节。通过建立可复现的实验记录和清晰的材料归档方式，"
            "系统不仅能够服务毕业设计验收，也能为后续同类通信抗干扰课题提供可继续迭代的基础。"
        ),
        (
            "因此，本课题后续工作的重点不应只放在单个指标提升上，还应围绕真实业务链路持续完善系统鲁棒性、可维护性和可复现实验能力，"
            "使平台在教学演示、算法验证和工程原型探索中都具有稳定价值。"
        ),
    ],
}


EXPANSIONS.setdefault("研究背景与意义", []).extend(
    [
        (
            "结合本人早期论文过程稿的论述，本课题的研究对象并不是单纯的离线干扰分类，而是面向视频业务的链路可靠性保障。"
            "视频业务对通信质量十分敏感，一旦干扰在毫秒级时间尺度上造成误码率上升，接收端就可能出现花屏、卡顿或帧丢失。"
            "因此，系统需要在业务质量明显恶化前尽早感知信道状态，并把干扰识别结果转化为可执行的传输参数调整。"
        ),
        (
            "在6G候选应用场景中，无人机、车联网、应急通信和密集异构网络会使频谱占用更加动态，通信链路面对的同频、邻频和恶意干扰也更复杂。"
            "这使得传统依靠固定阈值或固定抗干扰参数的方案难以长期保持稳定性能。本文围绕USRP软件无线电、深度学习识别和自适应视频传输三个层面展开，"
            "目的是构建一个能够从仿真走向实机验证的工程原型，而不仅是完成单项算法指标测试。"
        ),
    ]
)

EXPANSIONS.setdefault("功率谱特征提取", []).extend(
    [
        (
            "早期过程稿中将功率谱密度作为干扰检测的基础特征，这一点在最终系统中仍然保留。"
            "PSD能够直观反映信号能量在频域上的分布，尤其适合观察单音、多音、扫频和窄带类干扰引起的谱峰、带宽占用和能量扩散。"
            "因此，本文在CCNN七分类识别之外保留功率谱特征提取，用于可视化展示、干扰状态辅助判断和实验结果解释。"
        ),
        (
            "具体而言，谱峰值与噪底功率比可以衡量窄带强干扰的能量集中程度，频率重心可以反映干扰能量相对通信带宽中心的偏移，"
            "归一化谱熵能够描述功率谱分布是否趋于均匀，带外泄露功率比则用于观察邻频或越带干扰造成的频谱扩展。"
            "这些特征计算复杂度低、物理含义明确，与深度网络形成互补：前者便于解释和监控，后者负责完成更细粒度的类型识别。"
        ),
    ]
)

EXPANSIONS.setdefault("软件无线电与USRP B210", []).extend(
    [
        (
            "根据任务书、旧稿和USRP资料核对，USRP B210的射频收发范围覆盖70 MHz至6 GHz，支持最高56 MHz瞬时带宽，"
            "并通过USB 3.0接口与主机连接。UHD驱动为上层Python程序提供统一的设备发现、频率设置、增益控制和IQ数据收发接口。"
            "这些特性使其适合承担本课题从软件仿真到硬件在环验证的桥梁。"
        ),
        (
            "从系统集成角度看，USRP B210只负责射频前端和IQ采样转换，干扰注入、视频编解码、QPSK调制解调、CCNN推理和可视化仍由上位机软件完成。"
            "这种分工降低了硬件开发难度，也使毕业设计能够把主要精力集中在算法验证和链路闭环上。"
            "同时，真实硬件带来的频偏、采样时偏、增益波动和USB缓冲限制，也为第六章分析仿真与实机结果差异提供了依据。"
        ),
    ]
)

EXPANSIONS.setdefault("网络总体架构", []).extend(
    [
        (
            "旧稿中将识别网络概括为层次化的CCNN处理流程，这一思想在最终版本中进一步落实为卷积特征提取、残差学习、SE通道注意力、多头因果注意力和分类头。"
            "卷积模块负责捕获IQ序列中的局部幅相变化，残差结构保证较深网络的训练稳定性，SE模块对不同通道的重要性进行重标定，"
            "因果注意力则用于分析较长时间窗口内的干扰演化规律。"
        ),
        (
            "这种结构选择与外文文献中CNN和Transformer组合网络的思想一致，即先由卷积层提取符合信号物理特性的局部表示，"
            "再通过注意力机制补充全局依赖建模。与单纯堆叠卷积层相比，复合结构更适合处理扫频、多音和窄带调制等形态差异明显的干扰信号。"
        ),
    ]
)

EXPANSIONS.setdefault("实验环境", []).extend(
    [
        (
            "第六章实验数据与项目文件核对后，本文采用逐层验证口径：先使用离线生成样本检验CCNN七分类模型，再使用仿真模式验证视频TX/RX和自适应策略，"
            "最后通过USRP B210完成硬件识别、历史JSR扫描和120秒可视化长测。"
            "这种安排能够避免把临时短测结果误写成最终性能结论，也能更清楚地区分模型能力、仿真链路能力和实机链路能力。"
        ),
    ]
)


TRANSLATION_PARAGRAPHS = [
    "外文文献题名：A combination network of CNN and transformer for interference identification",
    "译文题名：用于干扰识别的CNN与Transformer组合网络",
    "原文作者：Hu Zhang, Meng Zhao, Min Zhang, Sheng Lin, Youqiang Dong, Hai Wang",
    "原文来源：Frontiers in Computational Neuroscience, 2023, 17:1309694.",
    (
        "通信干扰识别是电子对抗和智能通信系统中的关键问题。传统通信系统在受到有意或无意干扰时，通常只能根据接收功率、误码率或链路中断现象进行被动响应，"
        "难以及时判断干扰的具体类型。随着深度学习方法在信号处理领域的应用，研究者开始利用卷积神经网络、循环神经网络和Transformer等模型直接从信号数据中学习特征。"
        "然而，已有方法往往只强调局部特征或全局依赖中的某一方面，不能充分兼顾一维通信信号的局部波形变化和长距离上下文关系。针对这一不足，原文提出了一种将CNN与Transformer相结合的网络结构，"
        "通过联合利用卷积操作的局部特征提取能力和注意力机制的全局建模能力，提高通信干扰识别的准确性。"
    ),
    (
        "原文指出，干扰识别的目标是在缺少先验信息的情况下判断接收信号中干扰的类别，这对于抗干扰通信具有重要意义。"
        "现有干扰识别方法通常可以分为两类：一类是基于人工特征的方法，另一类是基于学习的方法。基于人工特征的方法会从接收信号中提取幅度、相位、小波变换结果或其他统计量，"
        "然后利用支持向量机等分类器进行判别。这类方法的优点是结构清晰、计算过程可解释，但其性能依赖特征设计，面对复杂干扰或信道变化时容易出现泛化能力不足的问题。"
        "基于学习的方法则利用深度神经网络自动学习特征表示，能够减少人工设计特征的工作量，并在调制识别和通信信号分类任务中取得较好效果。"
    ),
    (
        "在深度学习方法中，卷积神经网络常用于提取局部模式。对于通信信号而言，某些干扰类型会在短时间窗口或局部频谱范围内表现出明显特征，"
        "例如单音干扰形成稳定谱峰，线性扫频干扰表现为随时间变化的频率轨迹，多音干扰则包含多个离散频率成分。CNN通过局部卷积核能够捕获这些细粒度特征，"
        "并通过多层堆叠逐步形成更抽象的表示。可是，单纯依赖CNN也存在局限：当干扰特征跨越较长时间范围，或者需要比较不同位置之间的关系时，局部卷积感受野可能不足，"
        "网络需要更深层次或更复杂的结构才能覆盖全局信息。"
    ),
    (
        "Transformer结构的优势在于自注意力机制。自注意力可以直接计算序列中不同位置之间的相关性，因此适合处理具有长距离依赖的数据。"
        "在自然语言处理任务中，Transformer通过词嵌入表示离散词语，并利用多头注意力捕获上下文语义。原文认为，通信信号与自然语言存在明显差异："
        "一维信号并不天然具有词语边界，也难以直接套用词嵌入层来表示语义。因此，作者没有简单照搬标准Transformer，而是使用CNN替代词嵌入部分，"
        "先由卷积层从信号中提取更符合物理特性的局部表示，再送入Transformer编码结构进行全局关系建模。"
    ),
    (
        "基于这一思想，原文提出的CNNTF网络将CNN与Transformer进行组合。CNN部分负责处理原始信号数据，提取局部特征并形成适合后续编码器处理的表示；"
        "Transformer部分则利用注意力机制建立序列不同位置之间的依赖关系，使模型能够同时关注局部扰动和全局变化趋势。"
        "这种结构相较于简单拼接多个网络模块更具针对性，因为它根据通信信号的特点重新设计了Transformer的输入方式，避免了词嵌入对信号数据语义解释能力不足的问题。"
        "该设计思想对本文的视频传输平台也具有启发意义：干扰识别模型不仅要能够识别局部频谱异常，还要能够综合判断较长时间窗口内的变化规律。"
    ),
    (
        "原文进一步提出了结合交叉注意力机制的CNNTF-CA模型。通信干扰信号既可以在时域中观察，也可以在频域中观察。时域数据能够反映波形随时间变化的规律，"
        "频域数据能够反映能量在不同频率上的分布。传统做法往往需要人工进行时频分析，例如短时傅里叶变换或小波变换，再将不同域的特征输入不同分支网络。"
        "这种方法虽然有效，但流程较长，模型结构也更复杂。CNNTF-CA通过交叉注意力计算层融合时域和频域特征，使网络能够自动学习两个特征域之间的联系，"
        "从而在不依赖复杂专门时频分析步骤的情况下利用多域信息。"
    ),
    (
        "在信号模型方面，原文考虑了五种单一干扰信号，包括单音干扰、多音干扰、线性扫频干扰、部分带噪声干扰和噪声调频干扰。"
        "接收信号由通信信号、干扰信号和加性高斯白噪声叠加而成。作者同时从时域和频域描述信号，利用快速傅里叶变换将时域采样转换为频域表示，"
        "再提取幅度谱和相位谱。这样的建模方式说明，干扰识别并不是单纯的模式分类问题，而是与通信信号的调制方式、载波频率、初始相位、噪声环境和干扰参数密切相关。"
        "因此，识别网络必须具有足够强的特征提取能力，才能在不同信道条件下保持稳定性能。"
    ),
    (
        "实验结果表明，CNN与Transformer的组合能够提升干扰识别效果。CNN提供局部特征提取能力，Transformer提供全局依赖建模能力，交叉注意力机制则进一步增强了时域和频域信息的融合。"
        "原文的贡献主要体现在三个方面：第一，提出用CNN替代标准Transformer中的词嵌入层，使网络更适合通信信号数据；第二，引入交叉注意力机制，使模型能够同时利用不同域的特征；"
        "第三，通过实验验证所提方法相较已有方法具有更好的识别性能。该研究说明，在干扰识别任务中，合理融合不同网络结构比单纯增加网络深度更重要，"
        "网络设计应围绕信号数据的物理特征和任务需求展开。"
    ),
    (
        "结合本文课题，该外文文献对面向信号干扰检测与识别的视频传输平台具有直接参考价值。本文所使用的CCNN模型同样强调复合结构，"
        "通过卷积模块、残差模块、通道注意力和因果注意力机制共同提升识别能力。原文关于局部特征和全局特征互补的观点，为本文选择复合卷积网络而不是单一浅层分类器提供了理论依据。"
        "此外，原文对时域、频域和注意力融合的讨论，也说明干扰识别结果可以作为通信系统自适应调整的重要输入。"
        "因此，该文献不仅支撑了论文中深度学习干扰识别部分的技术路线，也为后续扩展时频多模态识别和跨层自适应传输提供了参考方向。"
    ),
    (
        "从方法实现角度看，原文的网络设计体现了“先局部、后全局”的处理思想。通信干扰信号通常既包含短时间内的突变特征，也包含跨采样窗口的演化规律。"
        "例如，单音干扰在局部频谱上表现为稳定峰值，多音干扰则由多个离散峰值共同构成；线性扫频干扰的关键特征不在单个采样点，而在频率随时间连续变化的轨迹。"
        "如果模型只关注局部卷积特征，就可能忽略不同时间片段之间的关联；如果模型只使用全局注意力，又可能弱化局部波形细节。"
        "因此，CNN与Transformer组合能够形成互补：CNN负责把原始信号转换为具有物理意义的局部表示，Transformer负责在更大范围内分析这些表示之间的关系。"
    ),
    (
        "原文对交叉注意力的使用也具有较强工程意义。通信系统中的同一段接收信号可以从多个角度观察，时域波形反映幅度和相位随时间的变化，"
        "频域表示反映能量在频率上的分布，二者并不是相互独立的信息。对于窄带调频干扰，频域能量集中程度和时域相位变化都可能包含分类线索；"
        "对于扫频干扰，单独观察瞬时频谱可能只能看到局部峰值，而结合时间变化才能确定其扫频特征。交叉注意力机制的作用正是建立不同特征域之间的对应关系，"
        "让模型自动学习哪些时域片段应与哪些频域结构共同参与判断。这种思想有助于减少人工规则，提高模型面对复杂信号时的适应能力。"
    ),
    (
        "在实验评价方面，原文强调通过对比实验验证所提网络的有效性。对于干扰识别研究而言，仅报告总体准确率并不足够，还需要关注不同类别之间的混淆情况、"
        "不同信噪比或干信比条件下的性能变化以及模型推理开销。原因在于实际通信系统通常运行在动态环境中，干扰强度、信道噪声和接收增益都会随时间变化。"
        "一个离线准确率很高但推理耗时较长的模型，未必适合部署到实时监测系统；一个在高信噪比下表现优秀但在弱干扰场景中失效的模型，"
        "也难以满足早期预警需求。因此，原文的研究思路提醒本文在实验章节中同时关注准确率、鲁棒性、参数规模和实时性。"
    ),
    (
        "该文献还说明了深度学习模型在通信干扰识别中的一个重要趋势：模型结构需要与信号处理知识结合，而不能脱离任务特点盲目套用通用网络。"
        "标准Transformer最初服务于自然语言任务，其输入通常是离散词元；通信信号则是连续采样序列，采样值之间存在明确的物理含义。"
        "原文用CNN替代词嵌入，实质上是把信号局部结构编码为更适合注意力处理的特征序列。本文的CCNN模型也遵循类似原则，"
        "没有直接把所有采样点交给分类器，而是通过卷积、残差和注意力模块逐步提取特征。这样的模型设计更符合IQ信号识别任务的本质。"
    ),
    (
        "对于本文的视频传输平台而言，干扰识别模型不是孤立模块，而是自适应传输控制的前端感知器。原文提出的组合网络提高了干扰类别判别能力，"
        "这意味着系统可以根据不同干扰类型采取差异化策略。例如，当识别到轻度窄带干扰时，平台可以适当提高FEC冗余并降低JPEG质量；"
        "当识别到更强的扫频或多音干扰时，平台可以进一步降低帧率以保证关键帧传输。由此可见，干扰识别精度会直接影响上层业务质量。"
        "原文研究为本文建立“识别结果-传输参数”之间的映射关系提供了理论支持，也说明智能通信系统应将信号感知、模型推理和业务控制协同设计。"
    ),
    (
        "原文的研究还体现了实验数据构造的重要性。深度学习模型并不能凭空获得对干扰信号的理解，其性能依赖训练样本是否覆盖足够多的信道条件和参数变化。"
        "如果训练集只包含少量固定功率或固定频率的干扰样本，模型可能只是记住了特定样本的局部形态，而无法在新的JSR或新的频偏条件下保持准确。"
        "因此，通信干扰识别任务通常需要在样本生成阶段引入多种随机因素，例如载波频率、初始相位、噪声功率、干扰带宽和调制参数。"
        "这一点与本文构建宽JSR数据集的思路一致，即通过增加样本多样性提升模型在复杂电磁环境中的泛化能力。"
    ),
    (
        "在实时系统部署方面，原文方法也提示研究者需要关注模型复杂度。Transformer具有较强的全局建模能力，但如果网络规模过大，推理延迟和计算资源消耗会限制其实时应用。"
        "通信监测系统往往需要连续处理高速采样数据，若每个推理窗口都消耗过长时间，识别结果将滞后于实际信道变化，无法及时指导抗干扰策略。"
        "因此，复合网络设计应在准确率和轻量化之间取得折中。本文的CCNN模型参数量较小，适合嵌入视频传输平台进行在线推理，"
        "这与原文追求高效干扰识别的目标具有一致性。"
    ),
    (
        "从论文写作角度，该文献可以支撑研究现状、模型设计和实验分析三个部分。在研究现状中，它说明了CNN、Transformer及注意力机制在通信干扰识别领域的发展方向；"
        "在模型设计中，它证明复合网络能够同时利用局部和全局特征；在实验分析中，它提醒本文不能只讨论某一类别准确率，而应结合不同干扰类型和不同JSR条件解释模型表现。"
        "因此，翻译该文献不仅满足外文资料阅读要求，也能够反向帮助完善毕业论文的技术论证，使论文的算法选择更有文献依据。"
    ),
    (
        "该文献的局限性也值得注意。原文主要关注干扰识别模型本身，实验重点在分类准确率提升，而对识别结果如何进一步驱动通信系统参数调整讨论较少。"
        "在真实业务系统中，干扰识别只是第一步，系统还需要根据识别结果决定是否降低码率、增加冗余、切换频点或改变调制方式。"
        "因此，本文在借鉴其模型思想的同时，将识别模块放入视频传输闭环中验证，使算法输出能够转化为具体的传输策略。"
        "这种扩展能够弥补单纯离线分类研究与工程应用之间的距离。"
    ),
    (
        "另外，原文关于多域特征融合的思想也为本文后续改进提供了方向。目前本文主要使用IQ序列和功率谱相关特征进行识别，"
        "若后续引入短时傅里叶变换图、连续小波变换图或星座图统计特征，则可以构建更完整的多模态干扰识别模型。"
        "不同模态的信息可以通过注意力机制或特征拼接方式融合，再由轻量分类头输出干扰类型和置信度。"
        "这种改进有望提升低JSR弱干扰场景下的识别稳定性，并为更复杂的混合干扰识别奠定基础。"
    ),
    (
        "综合比较原文方法与本文工作，可以看出二者关注点既有联系也有差异。原文重点回答“怎样提高干扰识别模型的分类能力”，"
        "本文则进一步回答“识别结果如何服务视频传输链路”。前者强调网络结构创新，后者强调系统集成和闭环验证。"
        "在毕业设计中，将这类外文研究成果转化为本课题的技术路线，需要结合硬件条件、程序实现难度和实验可复现性进行取舍。"
        "因此，本文没有完全复制原文模型，而是在其局部与全局特征融合思想启发下，设计了适合USRP视频传输平台的轻量级CCNN识别模块。"
    ),
    (
        "这篇外文资料还提醒我们，通信干扰识别研究必须同时面对算法准确性和系统实用性两个层面。算法准确性决定模型能否区分不同干扰，"
        "系统实用性则决定模型能否在有限计算资源和实时链路约束下发挥作用。本文将其思想吸收到毕业设计中，既关注CCNN模型在测试集上的准确率，"
        "也关注模型推理结果能否及时反馈到视频编码和冗余控制。这样的理解有助于避免论文只停留在理论算法描述，而能够体现通信工程专业毕业设计所要求的系统实现能力。"
    ),
    (
        "在后续研究中，还可以进一步结合原文提出的交叉注意力机制，将本文平台采集到的时域IQ数据、频域功率谱和视频链路质量指标共同建模。"
        "这样不仅可以识别干扰类型，还可以预测干扰对业务质量的影响程度，从而让自适应传输策略由规则驱动逐步发展为数据驱动。"
        "这一方向能够把外文文献中的模型创新与本文系统平台的工程优势更紧密地结合起来，也能为后续智能抗干扰通信实验提供更完整的研究基础和持续改进路径。"
    ),
    (
        "综上所述，原文提出的CNN与Transformer组合网络面向通信干扰识别任务，解决了单一网络结构难以同时捕获局部特征和全局依赖的问题。"
        "通过使用CNN替代词嵌入、引入Transformer编码器和交叉注意力机制，模型能够更充分地利用通信信号在时间域和频率域中的特征信息。"
        "该方法对复杂电磁环境中的干扰识别具有较高参考价值，也体现了智能通信系统中算法模型与信号处理知识结合的发展趋势。"
        "对于本文的毕业设计而言，该文献为CCNN模型设计、干扰类型识别、特征融合和系统自适应传输策略提供了重要的外文理论依据。"
    ),
]


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def chinese_count(text: str) -> int:
    return len(re.findall(r"[\u4e00-\u9fff]", text))


def ensure_output_dir() -> None:
    OUT.mkdir(exist_ok=True)
    for path in [
        OUT / "过程稿_李智杰.zip",
        OUT / "附录_李智杰.zip",
        OUT / "抽检材料_李智杰.zip",
    ]:
        if path.exists():
            path.unlink()


def set_run_font(run, size: float | None = None, bold: bool | None = None, east_asia: str = "宋体") -> None:
    run.font.name = "Times New Roman"
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), east_asia)


def set_paragraph_format(
    paragraph,
    size: float = 10.5,
    bold: bool | None = None,
    align: WD_ALIGN_PARAGRAPH | None = None,
    first_line: bool = False,
    east_asia: str = "宋体",
) -> None:
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = Pt(18)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(0.74) if first_line else Cm(0)
    if not paragraph.runs:
        paragraph.add_run("")
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, east_asia=east_asia)


def set_paragraph_base(
    paragraph,
    align: WD_ALIGN_PARAGRAPH | None = None,
    first_line_cm: float = 0,
    before_pt: float = 0,
    after_pt: float = 0,
    line_pt: float = 18,
) -> None:
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = Pt(line_pt)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.space_before = Pt(before_pt)
    paragraph.paragraph_format.space_after = Pt(after_pt)
    paragraph.paragraph_format.first_line_indent = Cm(first_line_cm)


def reset_paragraph_runs(paragraph, text: str = ""):
    paragraph.text = ""
    return paragraph.add_run(text)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = Pt(18)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for name, size, bold, align in [
        ("Heading 1", 16, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 12, True, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 10.5, True, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.paragraph_format.alignment = align
        style.paragraph_format.line_spacing = Pt(18)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.first_line_indent = Cm(0)

    if "Caption" in styles:
        style = styles["Caption"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(10.5)
        style.font.bold = False
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.line_spacing = Pt(18)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.75)
        header_p = section.header.paragraphs[0]
        header_p.text = TITLE
        header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(header_p, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
        footer_p = section.footer.paragraphs[0]
        footer_p.text = ""
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_field(footer_p)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, size=9)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    parent.remove(element)


def insert_paragraph_after(paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def set_text(paragraph, text: str) -> None:
    paragraph.text = text


def find_paragraph(doc: Document, predicate) -> Paragraph:
    for paragraph in doc.paragraphs:
        if predicate(paragraph):
            return paragraph
    raise ValueError("Paragraph not found")


def remove_between(start_p: Paragraph, end_p: Paragraph) -> None:
    removing = False
    for paragraph in list(start_p._parent.paragraphs):
        if paragraph._p is start_p._p:
            removing = True
            continue
        if paragraph._p is end_p._p:
            break
        if removing:
            delete_paragraph(paragraph)


def add_update_fields_setting(doc: Document) -> None:
    settings = doc.settings._element
    for existing in settings.findall(qn("w:updateFields")):
        settings.remove(existing)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_toc_field(paragraph: Paragraph) -> None:
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录域：请在 Word 中更新目录"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])
    set_run_font(run, size=10.5)


def rebuild_manual_toc(doc: Document) -> Paragraph:
    toc_h = find_paragraph(doc, lambda p: clean(p.text) == "目录")
    first_chapter = find_paragraph(doc, lambda p: p.style.name == "Heading 1" and clean(p.text).startswith("第一章"))
    remove_between(toc_h, first_chapter)

    entries: list[tuple[int, str]] = []
    for paragraph in doc.paragraphs:
        style = paragraph.style.name if paragraph.style else ""
        text = clean(paragraph.text)
        if not text or text in {"目录", "致谢"}:
            continue
        if style == "Heading 1":
            entries.append((1, text))
        elif style == "Heading 2":
            entries.append((2, text))
        elif style == "Heading 3":
            entries.append((3, text))

    cursor = toc_h
    for level, text in entries:
        indent = "    " * (level - 1)
        cursor = insert_paragraph_after(cursor, f"{indent}{text}")
        size = 10.5 if level == 1 else 10
        bold = level == 1
        set_paragraph_format(cursor, size=size, bold=bold, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    cursor.runs[-1].add_break(WD_BREAK.PAGE)
    return cursor


def add_section_break_after(paragraph: Paragraph, fmt: str, start: int) -> None:
    body = paragraph._parent._element
    last_sect = body.sectPr
    if last_sect is None:
        return
    section = deepcopy(last_sect)
    pg = section.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        section.append(pg)
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:sectPr"))
    if existing is not None:
        p_pr.remove(existing)
    p_pr.append(section)
    final_pg = last_sect.find(qn("w:pgNumType"))
    if final_pg is None:
        final_pg = OxmlElement("w:pgNumType")
        last_sect.append(final_pg)
    final_pg.set(qn("w:fmt"), "decimal")
    final_pg.set(qn("w:start"), "1")


def build_final_docx() -> None:
    shutil.copyfile(MAIN_DOCX, FINAL_DOCX)
    doc = Document(str(FINAL_DOCX))
    configure_styles(doc)
    configure_sections(doc)

    original_front = list(doc.paragraphs[:8])
    first = doc.paragraphs[0]
    cover_items = [
        "北京信息科技大学",
        "",
        "毕业设计（论文）",
        "",
        "",
        f"题    目：{TITLE}",
        "",
        "",
        f"学    院：{COLLEGE}",
        "",
        "",
        f"专    业：{MAJOR}",
        "",
        "",
        f"学生姓名：{STUDENT}        班级/学号：{CLASS_ID}",
        "",
        "",
        f"指导老师/督导老师：{ADVISOR}",
        "",
        "",
        f"起止时间：2026年2月23日 至 {DATE_TEXT}",
    ]
    inserted: list[Paragraph] = []
    for text in cover_items:
        p = first.insert_paragraph_before(text)
        set_paragraph_format(p, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, east_asia="楷体_GB2312")
        inserted.append(p)
    inserted[-1].runs[-1].add_break(WD_BREAK.PAGE)

    decl_items = [
        ("毕业设计（论文）原创性声明", 16, True, WD_ALIGN_PARAGRAPH.CENTER),
        (
            f"本人郑重声明：所呈交的毕业设计（论文），题目为《{TITLE}》，是本人在导师指导下，进行研究工作所取得的成果。"
            "尽我所知，除了文中特别加以标注的内容外，本毕业设计（论文）的研究成果不包含任何他人创作的、已公开发表或者没有公开发表的作品的内容。"
            "对本毕业设计（论文）所涉及的研究工作做出贡献的其他个人和集体，均已在文中以明确方式标明并表示了谢意。"
            "本毕业设计（论文）原创性声明的法律责任由本人承担。",
            10.5,
            False,
            WD_ALIGN_PARAGRAPH.JUSTIFY,
        ),
        ("作者签名：", 10.5, False, WD_ALIGN_PARAGRAPH.LEFT),
        (DATE_TEXT, 10.5, False, WD_ALIGN_PARAGRAPH.RIGHT),
        ("毕业设计（论文）版权授权声明", 16, True, WD_ALIGN_PARAGRAPH.CENTER),
        (
            "本人完全了解北京信息科技大学关于收集、保存、使用本科生毕业设计（论文）的要求，按照学校要求提交毕业设计（论文）的印刷本和电子版本。"
            "学校有权保留毕业设计（论文）并向相关机构送交论文的电子版和纸质版，允许论文被查阅和借阅，可以采用影印、缩印或扫描等复制手段保存、汇编毕业设计（论文）。"
            "学校有权适当复制、公布论文的全部或部分内容。",
            10.5,
            False,
            WD_ALIGN_PARAGRAPH.JUSTIFY,
        ),
        ("指导教师签名：                 作者签名：", 10.5, False, WD_ALIGN_PARAGRAPH.LEFT),
        (f"{DATE_TEXT}                 {DATE_TEXT}", 10.5, False, WD_ALIGN_PARAGRAPH.RIGHT),
    ]
    for text, size, bold, align in decl_items:
        p = first.insert_paragraph_before(text)
        set_paragraph_format(p, size=size, bold=bold, align=align, first_line=(align == WD_ALIGN_PARAGRAPH.JUSTIFY))
        inserted.append(p)
    decl_break = first.insert_paragraph_before("")
    decl_break.add_run().add_break(WD_BREAK.PAGE)

    for paragraph in original_front:
        delete_paragraph(paragraph)

    abstract_h = find_paragraph(doc, lambda p: clean(p.text).replace(" ", "") == "摘要")
    english_h = find_paragraph(doc, lambda p: clean(p.text).lower() == "abstract")
    toc_h = find_paragraph(doc, lambda p: clean(p.text) == "目录")
    remove_between(abstract_h, english_h)
    cursor = abstract_h
    for text in CN_ABSTRACT + [CN_KEYWORDS]:
        cursor = insert_paragraph_after(cursor, text)
        set_paragraph_format(cursor, first_line=not text.startswith("关键词"), align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    cursor.runs[-1].add_break(WD_BREAK.PAGE)

    remove_between(english_h, toc_h)
    cursor = english_h
    for text in EN_ABSTRACT + [EN_KEYWORDS]:
        cursor = insert_paragraph_after(cursor, text)
        set_paragraph_format(cursor, first_line=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    cursor.runs[-1].add_break(WD_BREAK.PAGE)

    first_chapter = find_paragraph(doc, lambda p: p.style.name == "Heading 1" and clean(p.text).startswith("第一章"))
    remove_between(toc_h, first_chapter)

    renumber_headings(doc)
    insert_expansions(doc)
    align_video_source_to_actual_usage(doc)
    replace_usrp_section_with_checked_material(doc)
    remove_acknowledgements(doc)
    normalize_paragraph_styles(doc)
    format_cover_and_declarations(doc)
    toc_last = rebuild_manual_toc(doc)
    add_section_break_after(toc_last, "upperRoman", 1)
    add_update_fields_setting(doc)
    doc.save(str(FINAL_DOCX))


def strip_heading_number(text: str) -> str:
    text = clean(text)
    text = re.sub(r"^\d+(?:\.\d+)*\s*", "", text)
    return text.strip()


def renumber_headings(doc: Document) -> None:
    chapter = 0
    h2 = 0
    h3 = 0
    in_ending = False
    ending_h2 = 0
    for paragraph in doc.paragraphs:
        text = clean(paragraph.text)
        if not text:
            continue
        style = paragraph.style.name
        if style == "Heading 1":
            if text.startswith("第七章") or "总结与展望" in text:
                paragraph.text = "结束语"
                chapter = 0
                in_ending = True
                ending_h2 = 0
            elif re.match(r"^第[一二三四五六七八九十]+章", text):
                chapter += 1
                h2 = 0
                h3 = 0
                in_ending = False
            elif text.startswith("参考文献") or text.startswith("致谢"):
                in_ending = False
                chapter = 0
            set_paragraph_format(paragraph, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
        elif style == "Heading 2":
            title = strip_heading_number(text)
            if in_ending:
                ending_h2 += 1
                paragraph.text = f"{ending_h2} {title}"
            elif chapter:
                h2 += 1
                h3 = 0
                paragraph.text = f"{chapter}.{h2} {title}"
            set_paragraph_format(paragraph, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
        elif style == "Heading 3":
            title = strip_heading_number(text)
            if in_ending:
                paragraph.text = title
            elif chapter and h2:
                h3 += 1
                paragraph.text = f"{chapter}.{h2}.{h3} {title}"
            set_paragraph_format(paragraph, size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)


def remove_acknowledgements(doc: Document) -> None:
    try:
        ack_h = find_paragraph(doc, lambda p: clean(p.text).startswith("致谢"))
    except ValueError:
        return
    deleting = False
    for paragraph in list(doc.paragraphs):
        if paragraph._p is ack_h._p:
            deleting = True
        if deleting:
            delete_paragraph(paragraph)


def insert_expansions(doc: Document) -> None:
    used: set[str] = set()
    in_body = False
    for paragraph in list(doc.paragraphs):
        text = clean(paragraph.text)
        if paragraph.style.name == "Heading 1" and text.startswith("第一章"):
            in_body = True
        elif paragraph.style.name == "Heading 1" and text.startswith("参考文献"):
            break
        if not in_body:
            continue
        for key, paragraphs in EXPANSIONS.items():
            if key in text and key not in used:
                cursor = paragraph
                for addition in paragraphs:
                    cursor = insert_paragraph_after(cursor, addition)
                    set_paragraph_format(cursor, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                used.add(key)
                break


def replace_usrp_section_with_checked_material(doc: Document) -> None:
    body = doc.element.body
    children = list(body.iterchildren())
    start_idx: int | None = None
    end_idx: int | None = None
    start_paragraph: Paragraph | None = None

    for idx, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, doc)
        style = paragraph.style.name if paragraph.style else ""
        text = clean(paragraph.text)
        if start_idx is None:
            if style == "Heading 2" and "USRP实机模式测试" in text:
                start_idx = idx
                start_paragraph = paragraph
        elif style in {"Heading 1", "Heading 2"}:
            end_idx = idx
            break

    if start_idx is None or start_paragraph is None:
        return
    if end_idx is None:
        end_idx = len(children)

    for child in children[start_idx + 1 : end_idx]:
        body.remove(child)

    checked_items = [
        (
            "p",
            "本节根据《第六章测试数据.md》、USRP目录资料和GNURadio测试记录重新核对实机结论。"
            "终稿采用硬件识别、历史JSR扫描和120秒实时可视化长测作为支撑，不把短突发无同步脚本的临时结果作为最终准确率结论。",
        ),
        ("h3", "6.5.1 硬件识别与配置核对"),
        (
            "p",
            "USRP硬件能够被UHD正常识别：设备型号为USRP B210，序列号为7MFTKFU，连接方式为USB 3.0，"
            "UHD版本为4.9.0.0，固件版本为FW 8.0，FPGA版本为FPGA 16.0，寄存器回环测试通过。"
            "实机测试工作频率设置为2.45 GHz，采样率为2 MHz，该配置与论文系统参数和测试数据记录一致。",
        ),
        (
            "p",
            "仿真模式30秒基本功能测试中，系统成功初始化视频传输链路、CCNN模型和可视化面板，发送帧数125，接收帧数125，"
            "干扰事件5次，策略切换5次，最终策略为“轻度干扰-降质保传”；CCNN处理样本数为20,871,856，检测到干扰次数124。"
            "该结果用于证明闭环流程能够稳定运行。",
        ),
        ("h3", "6.5.2 CCNN与自适应策略核对结果"),
        (
            "p",
            "CCNN模型推理测试使用gnuradio/test_ccnn.py加载七分类模型，对LFM、MTJ、NAM、NFM、STJ、SIN和Clean各生成10个测试样本，"
            "七类均为10/10正确，总体结果为70/70，离线生成样本总体准确率为100.00%。"
            "该结论与第六章测试数据和gnuradio/test_report.txt记录一致。",
        ),
        (
            "p",
            "自适应策略测试使用gnuradio/test_adaptive.py验证8帧干扰确认、3帧等级确认和12帧无干扰恢复逻辑。"
            "连续严重干扰和首次轻度干扰均在第10帧完成档位切换，已确认干扰后的等级变化只需连续3帧稳定等级确认，"
            "无干扰恢复在第12帧切回none档位。该结果支撑本文关于滞后切换机制的设计说明。",
        ),
        ("h3", "6.5.3 历史JSR扫描与实机识别现象"),
        (
            "p",
            "历史USRP JSR扫描配置为TX增益25 dB、RX增益70 dB、每个JSR点20次测试。结果表明，在8 dB至16 dB区间，"
            "LFM、MTJ和STJ识别较稳定，NFM在10 dB以上表现较好，NAM随JSR变化存在波动，SIN类在该批实测中未形成有效识别结果。"
            "因此，终稿将SIN实机表现写为后续模型和实机采样一致性优化方向，而不是夸大为全部类别均稳定识别。",
        ),
        ("h3", "6.5.4 USRP实时可视化长测"),
        (
            "p",
            "120秒USRP full模式长测采用run_system_with_viz.py --mode full --duration 120。"
            "原始full流程发送帧数483，接收处理块数609，CCNN处理样本数121,800,000，检测到干扰次数605，策略切换7次；"
            "优化后仪表盘流程发送帧数482，接收处理块数823，CCNN处理样本数164,600,000，检测到干扰次数816，策略切换18次。"
            "两种流程最终策略均为“轻度干扰-降质保传”。",
        ),
        (
            "p",
            "优化后流程在no_gui=True路径下跳过完整视频解调和JPEG解码，仅保留RF监测、PSD/IQ显示和CCNN推理，"
            "因此120秒内接收处理块数由609提升至823，CCNN处理样本数由121,800,000提升至164,600,000。"
            "测试过程中UHD控制台仍出现接收overflow标记，说明当前Python实现存在连续收流与每块CCNN推理同线程竞争问题。"
            "这一现象不代表模型离线准确率下降，而说明实机链路后续应通过采集/推理解耦、降低推理频率或引入更底层实现来提升实时性。",
        ),
    ]

    cursor = start_paragraph
    for kind, text in checked_items:
        cursor = insert_paragraph_after(cursor, text)
        if kind == "h3":
            cursor.style = "Heading 3"
            set_paragraph_format(cursor, size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
        else:
            set_paragraph_format(cursor, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)


def align_video_source_to_actual_usage(doc: Document) -> None:
    for paragraph in list(doc.paragraphs):
        text = clean(paragraph.text)
        if not text:
            continue
        if "发送端负责从摄像头或测试视频源获取图像帧" in text:
            paragraph.text = text.replace(
                "发送端负责从摄像头或测试视频源获取图像帧",
                "发送端实际以系统内置测试图案或视频文件作为输入源获取图像帧",
            )
            set_paragraph_format(paragraph, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif text.startswith("系统支持三种视频源"):
            paragraph.text = "系统实际测试采用内置OpenCV测试图案作为视频源，同时保留视频文件输入接口，便于后续替换真实业务视频。"
            set_paragraph_format(paragraph, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif text.startswith("1.") and "测试图案" in text:
            paragraph.text = "测试图案（test）：内置OpenCV测试图像，是本文仿真模式和USRP长测采用的视频源。"
            set_paragraph_format(paragraph, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif text.startswith("2.") and ("摄像头" in text or "camera" in text or "VideoCapture" in text):
            delete_paragraph(paragraph)
        elif text.startswith("摄像头") and ("camera" in text or "VideoCapture" in text):
            delete_paragraph(paragraph)


def set_mixed_cover_line(
    paragraph: Paragraph,
    segments: list[tuple[str, float, bool]],
    align: WD_ALIGN_PARAGRAPH | None = WD_ALIGN_PARAGRAPH.LEFT,
    first_line_cm: float = 0.99,
    add_page_break: bool = False,
) -> None:
    paragraph.text = ""
    paragraph.style = "Normal"
    set_paragraph_base(paragraph, align=align, first_line_cm=first_line_cm)
    for text, size, bold in segments:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold, east_asia="楷体_GB2312")
    if add_page_break:
        paragraph.runs[-1].add_break(WD_BREAK.PAGE)


def format_cover_and_declarations(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = clean(paragraph.text)
        if not text:
            set_paragraph_base(paragraph, first_line_cm=0)
            continue
        if text == "北京信息科技大学":
            run = reset_paragraph_runs(paragraph, text)
            set_paragraph_base(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0)
            set_run_font(run, size=26, bold=False, east_asia="楷体_GB2312")
        elif text == "毕业设计（论文）":
            run = reset_paragraph_runs(paragraph, text)
            set_paragraph_base(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0)
            set_run_font(run, size=42, bold=True, east_asia="楷体_GB2312")
        elif text.startswith("题") and "目：" in text:
            set_mixed_cover_line(
                paragraph,
                [("题    目：", 14, True), (TITLE, 12, False)],
            )
        elif text.startswith("学") and "院：" in text:
            set_mixed_cover_line(
                paragraph,
                [("学    院：", 14, True), (COLLEGE, 14, False)],
            )
        elif text.startswith("专") and "业：" in text:
            set_mixed_cover_line(
                paragraph,
                [("专    业：", 14, True), (MAJOR, 14, False)],
            )
        elif text.startswith("学生姓名："):
            set_mixed_cover_line(
                paragraph,
                [
                    ("学生姓名：", 14, True),
                    (STUDENT, 14, False),
                    ("        班级/学号：", 14, True),
                    (CLASS_ID, 14, False),
                ],
            )
        elif text.startswith("指导老师/督导老师：") or text.startswith("指导教师："):
            set_mixed_cover_line(
                paragraph,
                [("指导老师/督导老师：", 14, True), (f"   {ADVISOR}", 14, False)],
            )
        elif text.startswith("起止时间："):
            set_mixed_cover_line(
                paragraph,
                [("起止时间：", 14, True), (f"     2026年2月23日 至 {DATE_TEXT}", 14, False)],
                add_page_break=True,
            )
        elif text in {"毕业设计（论文）原创性声明", "毕业设计（论文）版权授权声明"}:
            run = reset_paragraph_runs(paragraph, text)
            set_paragraph_base(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0)
            set_run_font(run, size=16, bold=True, east_asia="宋体")
        elif text.startswith("本人郑重声明") or text.startswith("本人完全了解"):
            set_paragraph_format(paragraph, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True)
        elif text.startswith("作者签名") or text.startswith("指导教师签名"):
            set_paragraph_format(paragraph, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
        elif re.match(r"^2026年5月20日(?:\s+2026年5月20日)?$", text):
            set_paragraph_format(paragraph, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=False)
        elif text == "摘要":
            break


def normalize_paragraph_styles(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)

    for paragraph in doc.paragraphs:
        text = clean(paragraph.text)
        if not text:
            continue
        style = paragraph.style.name
        if style == "Heading 1":
            set_paragraph_format(paragraph, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
        elif style == "Heading 2":
            set_paragraph_format(paragraph, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
        elif style == "Heading 3":
            set_paragraph_format(paragraph, size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
        elif re.match(r"^[图表]\d+-\d+", text):
            paragraph.style = "Caption"
            set_paragraph_format(paragraph, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
        else:
            set_paragraph_format(paragraph, size=10.5, bold=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        try:
            table.style = "Table Grid"
        except Exception:
            pass
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    set_paragraph_format(paragraph, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)


def build_check_docx() -> None:
    shutil.copyfile(FINAL_DOCX, CHECK_DOCX)
    doc = Document(str(CHECK_DOCX))
    abstract_h = find_paragraph(doc, lambda p: clean(p.text).replace(" ", "") == "摘要")
    for paragraph in list(doc.paragraphs):
        if paragraph._p is abstract_h._p:
            break
        delete_paragraph(paragraph)
    abstract_h.insert_paragraph_before(TITLE)
    title_p = doc.paragraphs[0]
    title_p.style = "Heading 1"
    set_paragraph_format(title_p, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)

    try:
        toc_h = find_paragraph(doc, lambda p: clean(p.text) == "目录")
        first_chapter = find_paragraph(doc, lambda p: p.style.name == "Heading 1" and clean(p.text).startswith("第一章"))
        remove_between(toc_h, first_chapter)
        delete_paragraph(toc_h)
    except Exception:
        pass

    try:
        ack_h = find_paragraph(doc, lambda p: clean(p.text).startswith("致谢"))
        for paragraph in list(doc.paragraphs):
            if paragraph._p is ack_h._p:
                deleting = True
            if "deleting" in locals() and deleting:
                delete_paragraph(paragraph)
    except Exception:
        pass

    doc.save(str(CHECK_DOCX))


def create_translation_docx() -> None:
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    doc.add_paragraph("北京信息科技大学", style=None)
    for p in doc.paragraphs:
        set_paragraph_format(p, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, east_asia="楷体")
    title = doc.add_paragraph("外文文献译文")
    set_paragraph_format(title, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, east_asia="楷体")
    meta = [
        f"题 目：{TITLE}",
        f"学 院：{COLLEGE}",
        f"专 业：{MAJOR}",
        f"学生姓名：{STUDENT}  班级/学号：{CLASS_ID}",
        f"指导教师：{ADVISOR}",
    ]
    for line in meta:
        p = doc.add_paragraph(line)
        set_paragraph_format(p, size=12, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False, east_asia="楷体")
    doc.add_page_break()
    for index, text in enumerate(TRANSLATION_PARAGRAPHS):
        p = doc.add_paragraph(text)
        if index < 4:
            set_paragraph_format(p, size=10.5, bold=(index == 1), align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
        else:
            set_paragraph_format(p, size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True)
    count_p = doc.add_paragraph(f"译文汉字数统计：{chinese_count(''.join(TRANSLATION_PARAGRAPHS))} 字")
    set_paragraph_format(count_p, size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
    doc.save(str(TRANSLATION_DOCX))


def create_index_docx(path: Path, title: str, rows: list[tuple[str, str]]) -> None:
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    p = doc.add_paragraph("北京信息科技大学")
    set_paragraph_format(p, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, east_asia="楷体")
    p = doc.add_paragraph(title)
    set_paragraph_format(p, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, east_asia="楷体")
    for line in [
        f"题 目：{TITLE}",
        f"学 院：{COLLEGE}",
        f"专 业：{MAJOR}",
        f"学生姓名：{STUDENT}  班级/学号：{CLASS_ID}",
        f"指导教师：{ADVISOR}",
    ]:
        p = doc.add_paragraph(line)
        set_paragraph_format(p, size=12, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False, east_asia="楷体")
    doc.add_page_break()
    p = doc.add_paragraph("材料目录")
    p.style = "Heading 1"
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "序号"
    hdr[1].text = "材料"
    hdr[2].text = "说明"
    for idx, (name, note) in enumerate(rows, 1):
        cells = table.add_row().cells
        cells[0].text = str(idx)
        cells[1].text = name
        cells[2].text = note
    normalize_paragraph_styles(doc)
    doc.save(str(path))


def build_index_docs() -> None:
    create_index_docx(
        APPENDIX_INDEX_DOCX,
        "毕业设计（论文）附录",
        [
            ("开题报告", "使用已有李智杰开题报告。"),
            ("设计图纸/系统架构图", "包含系统图片、USRP原理图及实验相关图件。"),
            ("计算机程序", "包含CCNN、视频传输、USRP测试等核心源码。"),
            ("重要原始数据与测试说明", "包含第六章测试数据、实验日志和结果说明。"),
            ("外文文献译文", "译文不少于5000汉字。"),
            ("外文文献原文", "包含Frontiers外文原文PDF。"),
        ],
    )
    create_index_docx(
        INSPECTION_INDEX_DOCX,
        "本科毕业论文抽检支撑材料清单",
        [
            ("开题报告与进展材料", "支撑选题、研究过程和阶段性检查。"),
            ("核心程序", "支撑系统实现、模型训练推理和USRP链路测试。"),
            ("测试数据与结果", "支撑论文第六章实验结果。"),
            ("系统说明文档", "支撑平台功能、运行方式和工程结构。"),
            ("外文原文与译文", "支撑外文资料阅读与翻译要求。"),
            ("论文与材料核对报告", "说明终稿与旧稿、测试数据、程序说明和过程材料的对应关系。"),
        ],
    )


def create_alignment_report_docx() -> None:
    create_index_docx(
        ALIGNMENT_REPORT_DOCX,
        "论文与毕设材料核对报告",
        [
            (
                "桌面旧稿",
                f"已参考 {DESKTOP_DRAFT_DOCX}。旧稿中的6G复杂电磁环境、视频业务敏感性、功率谱特征、USRP平台说明和CCNN层次化结构已融合到终稿相关章节。",
            ),
            (
                "毕设文件夹当前论文",
                f"以 {MAIN_DOCX} 作为正文主体基准，保留七章结构、图表、实验章节和参考文献，再按工作手册调整终稿装订口径。",
            ),
            (
                "第六章测试数据",
                "已按第六章测试数据.md核对并重写USRP实机测试表述：采用30秒仿真闭环、CCNN 70/70推理测试、自适应策略测试、历史JSR扫描和120秒USRP长测作为最终证据。",
            ),
            (
                "任务书、开题报告和进展报告",
                "已放入过程稿、附录和抽检材料包；系统终稿Word/PDF不重复塞入任务书、开题报告、评语表或成绩表。",
            ),
            (
                "程序与工程说明",
                "已核对并归档gnuradio、CCNN、USRP、README、测试数据摘要和系统说明文档，抽检材料排除了虚拟环境、缓存和大型重复目录。",
            ),
            (
                "外文原文与译文",
                "外文原文使用fncom-17-1309694.pdf，译文文件为外文文献译文_李智杰.docx，译文汉字数不少于5000。",
            ),
        ],
    )


def add_file(zf: zipfile.ZipFile, path: Path, arcname: str) -> None:
    if path.exists() and path.is_file() and not path.name.startswith("~$"):
        zf.write(path, arcname)


def add_tree(
    zf: zipfile.ZipFile,
    base: Path,
    arc_base: str,
    extensions: set[str],
    exclude_parts: set[str] | None = None,
    max_file_mb: float = 50,
) -> None:
    if not base.exists():
        return
    exclude_parts = exclude_parts or set()
    for path in base.rglob("*"):
        if not path.is_file() or path.name.startswith("~$"):
            continue
        if any(part in exclude_parts for part in path.parts):
            continue
        if path.suffix.lower() not in extensions:
            continue
        if path.stat().st_size > max_file_mb * 1024 * 1024:
            continue
        relative = path.relative_to(base)
        zf.write(path, str(Path(arc_base) / relative))


def build_zips() -> None:
    progress_zip = OUT / "过程稿_李智杰.zip"
    appendix_zip = OUT / "附录_李智杰.zip"
    inspection_zip = OUT / "抽检材料_李智杰.zip"
    progress_files = [
        ROOT / "进展" / "李智杰毕业设计任务书.docx",
        ROOT / "进展" / "李智杰开题报告.docx",
        ROOT / "进展" / "毕业设计进展报告.docx",
        ROOT / "进展" / "毕业设计综合进展报告.docx",
        ROOT / "进展" / "本科毕设进展.docx",
        ROOT / "进展" / "CCNN干扰信号识别.docx",
        DESKTOP_DRAFT_DOCX,
        ROOT / "毕业论文_李智杰_备份_20260415.docx",
        ROOT / "毕业论文_李智杰_自动备份_20260507_135744.docx",
        ROOT / "毕业论文_李智杰_自动备份_20260507_143531.docx",
        ROOT / "答辩PPT_李智杰.pptx",
        ROOT / "毕业论文.md",
        ROOT / "第六章测试数据.md",
        ALIGNMENT_REPORT_DOCX,
    ]
    with zipfile.ZipFile(progress_zip, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in progress_files:
            add_file(zf, path, f"过程稿/{path.name}")
        add_tree(zf, ROOT / "codex_prompts", "过程记录/codex_prompts", {".md"})

    common_files = [
        APPENDIX_INDEX_DOCX,
        INSPECTION_INDEX_DOCX,
        ALIGNMENT_REPORT_DOCX,
        ROOT / "进展" / "李智杰开题报告.docx",
        ROOT / "第六章测试数据.md",
        ROOT / "MATLAB_vs_Python.md",
        ROOT / "README.md",
        TRANSLATION_DOCX,
        FOREIGN_PDF,
        ROOT / "figure" / "微信图片_20260326204316_346_85.png",
        ROOT / "USRP" / "B210mini" / "原理图" / "miniRev0.pdf",
        ROOT / "USRP" / "usrp_jsr_accuracy.pdf",
        ROOT / "gnuradio" / "usrp_jsr_accuracy.pdf",
    ]
    for zip_path, prefix in [(appendix_zip, "附录"), (inspection_zip, "抽检材料")]:
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            seen_names: set[str] = set()
            for path in common_files:
                arc_name = path.name
                if arc_name in seen_names:
                    arc_name = f"{path.parent.name}_{path.name}"
                seen_names.add(arc_name)
                add_file(zf, path, f"{prefix}/{arc_name}")
            add_tree(
                zf,
                ROOT / "gnuradio",
                f"{prefix}/程序/gnuradio",
                {".py", ".md", ".txt", ".json", ".pdf", ".sh"},
                exclude_parts={"__pycache__", ".git"},
                max_file_mb=20,
            )
            add_tree(
                zf,
                ROOT / "CCNN" / "3_scripts",
                f"{prefix}/程序/CCNN/3_scripts",
                {".py", ".md", ".txt", ".m"},
                max_file_mb=20,
            )
            add_tree(
                zf,
                ROOT / "CCNN" / "5_docs",
                f"{prefix}/程序/CCNN/5_docs",
                {".md", ".txt", ".png"},
                max_file_mb=20,
            )
            add_tree(
                zf,
                ROOT / "CCNN" / "2_models",
                f"{prefix}/程序/CCNN/2_models",
                {".md", ".txt", ".pth", ".onnx", ".om"},
                max_file_mb=20,
            )
            add_tree(
                zf,
                ROOT / "参考文献",
                f"{prefix}/参考文献",
                {".docx", ".pdf"},
                max_file_mb=20,
            )
            add_tree(
                zf,
                ROOT / "USRP",
                f"{prefix}/程序/USRP",
                {".py", ".json", ".pdf", ".docx", ".txt"},
                exclude_parts={"用户用到得软件", "libusb-1.0.27", "Vivado工程"},
                max_file_mb=30,
            )


def collect_stats() -> dict[str, int | float]:
    doc = Document(str(FINAL_DOCX))
    paragraphs = [(p.style.name if p.style else "", clean(p.text)) for p in doc.paragraphs if clean(p.text)]
    first_chapter = next(i for i, (style, text) in enumerate(paragraphs) if style == "Heading 1" and text.startswith("第一章"))
    refs = next(i for i, (style, text) in enumerate(paragraphs) if style == "Heading 1" and text.startswith("参考文献"))
    body_text = "\n".join(text for _, text in paragraphs[first_chapter:refs])
    translation_text = "\n".join(clean(p.text) for p in Document(str(TRANSLATION_DOCX)).paragraphs)
    return {
        "body_chinese_chars": chinese_count(body_text),
        "translation_chinese_chars": chinese_count(translation_text),
        "process_zip_mb": round((OUT / "过程稿_李智杰.zip").stat().st_size / 1024 / 1024, 2),
        "appendix_zip_mb": round((OUT / "附录_李智杰.zip").stat().st_size / 1024 / 1024, 2),
        "inspection_zip_mb": round((OUT / "抽检材料_李智杰.zip").stat().st_size / 1024 / 1024, 2),
    }


def main() -> None:
    ensure_output_dir()
    build_final_docx()
    create_translation_docx()
    build_index_docs()
    build_check_docx()
    create_alignment_report_docx()
    build_zips()
    stats = collect_stats()
    for key, value in stats.items():
        print(f"{key}: {value}")
    print(f"output_dir: {OUT}")


if __name__ == "__main__":
    main()
