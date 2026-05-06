#!/usr/bin/env python3
"""
将《毕业论文.md》转换为学校格式要求的 Word 文档。

实现范围：
- Markdown 标题、正文、加粗、行内代码、链接文本
- 独立公式块、代码块
- Markdown 表格
- Markdown 图片引用
- 页眉、页脚页码、A4页面和正文/标题样式
"""
from __future__ import annotations

import argparse
import os
import re
import shutil
from datetime import datetime
from pathlib import Path

import subprocess
import tempfile
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree


PROJECT_ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = PROJECT_ROOT / "毕业论文.md"
DEFAULT_OUTPUT = PROJECT_ROOT / "毕业论文_李智杰.docx"
HEADER_TEXT = "北京信息科技大学本科毕业设计（论文）"


IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]+\)")
TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


class ConvertStats:
    def __init__(self) -> None:
        self.paragraphs = 0
        self.tables = 0
        self.images = 0
        self.formulas = 0
        self.missing_images: list[str] = []


def set_run_font(run, font_name="宋体", size=12, bold=None, italic=None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_format(paragraph, align=None, first_line=True) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(6)
    if first_line:
        fmt.first_line_indent = Cm(0.74)
    if align is not None:
        paragraph.alignment = align


def set_cell_text(cell, text: str, bold=False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, first_line=False)
    add_inline_runs(p, text, bold=bold)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    set_run_font(run, size=10.5)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    header = section.header.paragraphs[0]
    header.text = HEADER_TEXT
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        set_run_font(run, size=10.5)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    style_map = {
        "Heading 1": ("黑体", 16, True, WD_ALIGN_PARAGRAPH.CENTER),
        "Heading 2": ("黑体", 14, True, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 3": ("黑体", 12, True, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 4": ("黑体", 12, True, WD_ALIGN_PARAGRAPH.LEFT),
    }
    for style_name, (font_name, size, bold, align) in style_map.items():
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.paragraph_format.alignment = align
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def normalize_inline(text: str) -> str:
    text = LINK_RE.sub(r"\1", text)
    return text.replace("<br>", "\n")


def add_inline_runs(paragraph, text: str, bold=False) -> None:
    text = normalize_inline(text)
    pos = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, bold=bold)
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, font_name="Consolas", size=11)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, bold=bold)


def clean_heading_text(text: str) -> str:
    return re.sub(r"^\s*\d+[\.\s]+", "", text).strip()


def add_heading(document: Document, level: int, text: str) -> None:
    text = clean_heading_text(text)
    if level == 1:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(p, WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
        run = p.add_run(text)
        set_run_font(run, font_name="黑体", size=20, bold=True)
        return

    style_level = min(max(level - 1, 1), 4)
    p = document.add_paragraph(style=f"Heading {style_level}")
    p.text = ""
    add_inline_runs(p, text, bold=True)
    p.paragraph_format.first_line_indent = None
    if style_level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def is_table_separator(row: str) -> bool:
    cells = [c.strip() for c in row.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        if is_table_separator(line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def add_table(document: Document, rows: list[list[str]], stats: ConvertStats) -> None:
    if not rows:
        return
    max_cols = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=max_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows):
        for c_idx in range(max_cols):
            text = row[c_idx] if c_idx < len(row) else ""
            set_cell_text(table.cell(r_idx, c_idx), text, bold=(r_idx == 0))
    stats.tables += 1


def add_image(document: Document, alt: str, path_text: str, stats: ConvertStats) -> None:
    image_path = (PROJECT_ROOT / path_text.strip()).resolve()
    if not image_path.exists():
        stats.missing_images.append(path_text.strip())
        p = document.add_paragraph()
        set_paragraph_format(p)
        add_inline_runs(p, f"[图片缺失：{path_text.strip()}]")
        return

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(13.5))
    stats.images += 1

    if alt:
        cap = document.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(cap, WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
        add_inline_runs(cap, alt)
        for run in cap.runs:
            set_run_font(run, size=10.5)


def add_paragraph(document: Document, text: str, stats: ConvertStats) -> None:
    stripped = text.strip()
    if not stripped:
        return

    if stripped.startswith("- "):
        p = document.add_paragraph(style="List Bullet")
        set_paragraph_format(p, first_line=False)
        add_inline_runs(p, stripped[2:])
    elif re.match(r"^\d+\.\s+", stripped):
        p = document.add_paragraph(style="List Number")
        set_paragraph_format(p, first_line=False)
        add_inline_runs(p, re.sub(r"^\d+\.\s+", "", stripped))
    else:
        p = document.add_paragraph()
        set_paragraph_format(p)
        add_inline_runs(p, stripped)
    stats.paragraphs += 1


OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
OMML_NS_PREFIX = f"{{{OMML_NS}}}"


def build_formula_map(lines: list[str]) -> dict[int, etree._Element]:
    """Pre-convert all $$...$$ formulas to OMML using pandoc.

    Returns a dict mapping line-number → OMML element.
    """
    formula_lines: list[tuple[int, str]] = []
    in_formula = False
    formula_buf: list[str] = []
    formula_start = -1

    for idx, line in enumerate(lines):
        stripped = line.strip()
        if stripped == "$$":
            if in_formula:
                latex = "\n".join(formula_buf).strip()
                if latex:
                    formula_lines.append((formula_start, latex))
                in_formula = False
                formula_buf = []
            else:
                in_formula = True
                formula_buf = []
                formula_start = idx
            continue
        if in_formula:
            formula_buf.append(line)
            continue
        # Single-line $$...$$
        if stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) > 4:
            latex = stripped[2:-2].strip()
            if latex:
                formula_lines.append((idx, latex))

    if not formula_lines:
        return {}

    # Build one pandoc input with anchors for each formula
    pandoc_parts = []
    for i, (_, latex) in enumerate(formula_lines):
        pandoc_parts.append(f"$FORMULA_START_{i}$")
        pandoc_parts.append(f"$$\n{latex}\n$$")
        pandoc_parts.append(f"$FORMULA_END_{i}$")
    pandoc_input = "\n".join(pandoc_parts)

    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "input.md"
            output_path = Path(tmpdir) / "output.docx"
            input_path.write_text(pandoc_input, encoding="utf-8")
            result = subprocess.run(
                ["pandoc", str(input_path), "-o", str(output_path), "--to", "docx"],
                capture_output=True, text=True, timeout=30,
            )
            if result.returncode != 0:
                print(f"  [公式] pandoc 转换失败: {result.stderr[:200]}")
                return {}

            with zipfile.ZipFile(output_path) as zf:
                doc_bytes = zf.read("word/document.xml")

        tree = etree.fromstring(doc_bytes)
        # Collect all oMathPara and oMath elements
        omml_elements = tree.findall(f".//{OMML_NS_PREFIX}oMathPara")
        if not omml_elements:
            omml_elements = tree.findall(f".//{OMML_NS_PREFIX}oMath")

        if len(omml_elements) != len(formula_lines):
            print(f"  [公式] OMML数量({len(omml_elements)}) ≠ 公式数量({len(formula_lines)})，跳过公式渲染")
            return {}

        result_map: dict[int, etree._Element] = {}
        for (line_idx, _), omml in zip(formula_lines, omml_elements):
            result_map[line_idx] = omml
        print(f"  [公式] 成功转换 {len(result_map)} 个公式为 Word 原生格式")
        return result_map

    except Exception as e:
        print(f"  [公式] 转换异常，将使用纯文本: {e}")
        return {}


def add_formula_paragraph(document: Document, omml_element: etree._Element | None,
                          fallback_text: str) -> None:
    """Add a formula paragraph: OMML if available, else plain text."""
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    if omml_element is not None:
        p._element.append(omml_element)
    else:
        run = p.add_run(fallback_text)
        set_run_font(run, size=12)


def convert_markdown(input_path: Path, output_path: Path, backup=True) -> ConvertStats:
    stats = ConvertStats()
    lines = input_path.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)

    # Pre-convert all formulas to OMML using pandoc
    print("正在转换公式...")
    formula_map = build_formula_map(lines)

    i = 0
    in_code = False
    code_lines: list[str] = []
    in_formula = False
    formula_lines: list[str] = []
    formula_start_line = -1

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                p = document.add_paragraph()
                set_paragraph_format(p, first_line=False)
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, font_name="Consolas", size=10)
                in_code = False
                code_lines = []
            else:
                in_code = True
                code_lines = []
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) > 4:
            add_formula_paragraph(document, formula_map.get(i), stripped)
            stats.formulas += 1
            i += 1
            continue

        if stripped == "$$":
            if in_formula:
                fallback = "$$" + "\n".join(formula_lines) + "$$"
                add_formula_paragraph(document, formula_map.get(formula_start_line), fallback)
                stats.formulas += 1
                in_formula = False
                formula_lines = []
            else:
                in_formula = True
                formula_lines = []
                formula_start_line = i
            i += 1
            continue
        if in_formula:
            formula_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue
        if stripped == "---":
            i += 1
            continue

        heading = HEADING_RE.match(line)
        if heading:
            add_heading(document, len(heading.group(1)), heading.group(2))
            i += 1
            continue

        image = IMAGE_RE.search(line)
        if image:
            add_image(document, image.group(1), image.group(2), stats)
            i += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(document, parse_table(table_lines), stats)
            continue

        add_paragraph(document, line, stats)
        i += 1

    if output_path.exists() and backup:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = output_path.with_name(f"{output_path.stem}_自动备份_{stamp}{output_path.suffix}")
        shutil.copy2(output_path, backup_path)
        print(f"已备份旧文件：{backup_path}")

    document.save(output_path)
    return stats


def main() -> None:
    parser = argparse.ArgumentParser(description="生成毕业论文Word文档")
    parser.add_argument("--input", default=str(DEFAULT_INPUT), help="Markdown源文件")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT), help="Word输出文件")
    parser.add_argument("--no-backup", action="store_true", help="不备份已有输出文件")
    args = parser.parse_args()

    input_path = Path(args.input).resolve()
    output_path = Path(args.output).resolve()
    if not input_path.exists():
        raise FileNotFoundError(f"源文件不存在：{input_path}")

    stats = convert_markdown(input_path, output_path, backup=not args.no_backup)
    print(f"已生成：{output_path}")
    print(f"段落：{stats.paragraphs}，表格：{stats.tables}，图片：{stats.images}，公式：{stats.formulas}")
    if stats.missing_images:
        print("缺失图片：")
        for item in stats.missing_images:
            print(f"  - {item}")


if __name__ == "__main__":
    main()
